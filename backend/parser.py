"""Parse bank statement CSV / Excel / PDF into a common transaction list."""

from __future__ import annotations

import csv
import io
import re
from datetime import datetime
from typing import Any

from dateutil import parser as dateparser
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DATE_KEYS = ("date", "txn date", "value date", "transaction date", "posting date", "tran date")
DESC_KEYS = (
    "description",
    "narration",
    "particulars",
    "details",
    "remarks",
    "transaction remarks",
    "narrative",
)
DEBIT_KEYS = ("debit", "withdrawal", "withdrawals", "dr", "debit amount", "withdrawal amt.")
CREDIT_KEYS = ("credit", "deposit", "deposits", "cr", "credit amount", "deposit amt.")
AMOUNT_KEYS = ("amount", "txn amount", "transaction amount")
BAL_KEYS = ("balance", "closing balance", "running balance", "available balance")
REF_KEYS = ("ref", "cheque", "chq", "instrument", "utr", "reference", "txn id", "transaction id")


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").strip().lower())


def _parse_amount(val: Any) -> float | None:
    if val is None or val == "":
        return None
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip()
    if not s or s in {"-", "—", "NA", "N/A"}:
        return None
    neg = False
    if s.startswith("(") and s.endswith(")"):
        neg = True
        s = s[1:-1]
    s = s.replace(",", "").replace("₹", "").replace("Rs.", "").replace("INR", "").strip()
    s = re.sub(r"[^\d.\-]", "", s)
    if not s or s in {".", "-"}:
        return None
    try:
        n = float(s)
        return -n if neg else n
    except ValueError:
        return None


def _parse_date(val: Any) -> str:
    if val is None or val == "":
        return ""
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d")
    s = str(val).strip()
    if not s:
        return ""
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%b-%Y", "%d %b %Y", "%d-%m-%y", "%d/%m/%y"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    try:
        return dateparser.parse(s, dayfirst=True).strftime("%Y-%m-%d")
    except (ValueError, OverflowError, TypeError):
        return s


def _pick_col(headers: list[str], keys: tuple[str, ...]) -> int | None:
    norms = [_norm(h) for h in headers]
    for i, h in enumerate(norms):
        for k in keys:
            if k == h or k in h:
                return i
    return None


def _rows_from_csv(raw: bytes) -> list[list[str]]:
    text = raw.decode("utf-8-sig", errors="replace")
    sample = text[:4000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    return [list(r) for r in reader]


def _rows_from_xlsx(raw: bytes) -> list[list[str]]:
    wb = load_workbook(io.BytesIO(raw), data_only=True)
    ws = wb.active
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append(["" if c is None else c for c in row])
    return rows


def _rows_from_pdf(raw: bytes) -> list[list[str]]:
    import pdfplumber

    rows: list[list[str]] = []
    with pdfplumber.open(io.BytesIO(raw)) as pdf:
        for page in pdf.pages:
            tables = page.extract_tables() or []
            for table in tables:
                for r in table:
                    rows.append([("" if c is None else str(c).replace("\n", " ")) for c in r])
            if not tables:
                text = page.extract_text() or ""
                for line in text.splitlines():
                    parts = re.split(r"\s{2,}", line.strip())
                    if len(parts) >= 3:
                        rows.append(parts)
    return rows


def _find_header(rows: list[list[Any]]) -> tuple[int, list[str]]:
    best_i, best_score = 0, -1
    for i, row in enumerate(rows[:40]):
        joined = " ".join(_norm(str(c)) for c in row)
        score = 0
        if any(k in joined for k in DATE_KEYS):
            score += 3
        if any(k in joined for k in DESC_KEYS):
            score += 2
        if any(k in joined for k in (*DEBIT_KEYS, *CREDIT_KEYS, *AMOUNT_KEYS)):
            score += 3
        if score > best_score:
            best_score, best_i = score, i
    headers = [str(c) if c is not None else "" for c in rows[best_i]]
    return best_i, headers


def parse_statement(filename: str, raw: bytes) -> dict:
    lower = filename.lower()
    if lower.endswith(".pdf"):
        rows = _rows_from_pdf(raw)
        source = "pdf"
    elif lower.endswith((".xlsx", ".xls")):
        rows = _rows_from_xlsx(raw)
        source = "excel"
    else:
        rows = _rows_from_csv(raw)
        source = "csv"

    if not rows:
        return {"transactions": [], "source": source, "warning": "No rows found in file."}

    header_i, headers = _find_header(rows)
    date_i = _pick_col(headers, DATE_KEYS)
    desc_i = _pick_col(headers, DESC_KEYS)
    debit_i = _pick_col(headers, DEBIT_KEYS)
    credit_i = _pick_col(headers, CREDIT_KEYS)
    amount_i = _pick_col(headers, AMOUNT_KEYS)
    bal_i = _pick_col(headers, BAL_KEYS)
    ref_i = _pick_col(headers, REF_KEYS)

    txns = []
    for raw_row in rows[header_i + 1 :]:
        cells = list(raw_row) + [""] * 12
        date = _parse_date(cells[date_i]) if date_i is not None else ""
        desc = str(cells[desc_i]).strip() if desc_i is not None else " ".join(str(c) for c in raw_row)
        debit = _parse_amount(cells[debit_i]) if debit_i is not None else None
        credit = _parse_amount(cells[credit_i]) if credit_i is not None else None
        amount = _parse_amount(cells[amount_i]) if amount_i is not None else None
        balance = _parse_amount(cells[bal_i]) if bal_i is not None else None
        ref = str(cells[ref_i]).strip() if ref_i is not None else ""

        if debit is None and credit is None and amount is not None:
            if amount < 0:
                debit, credit = abs(amount), None
            else:
                credit, debit = amount, None

        if not date and not desc:
            continue
        if debit is None and credit is None:
            continue

        cr = float(credit or 0)
        dr = float(debit or 0)
        signed = cr - dr
        txns.append(
            {
                "date": date,
                "description": re.sub(r"\s+", " ", desc)[:400],
                "reference": ref,
                "debit": dr if dr else None,
                "credit": cr if cr else None,
                "amount": signed,
                "balance": balance,
                "type": "receipt" if signed > 0 else "payment",
                "suggested_ledger": _suggest_ledger(desc),
            }
        )

    return {
        "transactions": txns,
        "source": source,
        "headers": headers,
        "count": len(txns),
    }


def _suggest_ledger(desc: str) -> str:
    d = desc.upper()
    rules = [
        (("UPI", "GPAY", "PHONEPE", "PAYTM"), "Sundry Debtors"),
        (("NEFT", "RTGS", "IMPS"), "Sundry Debtors"),
        (("SALARY", "PAYROLL"), "Salary"),
        (("GST", "TAX"), "GST Payable"),
        (("INT.", "INTEREST", "INT PAID", "INT CR"), "Bank Interest"),
        (("CHQ", "CHEQUE"), "Sundry Creditors"),
        (("ATM", "CASH WDL", "CASH WITHDRAW"), "Cash"),
        (("POS", "CARD"), "Bank Charges"),
        (("CHARGE", "FEE", "SMS"), "Bank Charges"),
        (("RENT",), "Rent"),
        (("ELECTRIC", "BSEB", "TORRENT", "UGVCL"), "Electricity"),
    ]
    for keys, ledger in rules:
        if any(k in d for k in keys):
            return ledger
    return "Suspense"


def transactions_to_xlsx(txns: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Bank Statement"

    headers = [
        "Date",
        "Description",
        "Reference",
        "Debit",
        "Credit",
        "Amount",
        "Balance",
        "Type",
        "Suggested Tally Ledger",
    ]
    header_fill = PatternFill("solid", fgColor="0F3D3E")
    header_font = Font(color="F4EFE6", bold=True, name="Calibri", size=11)
    thin = Border(
        left=Side(style="thin", color="D6CFC2"),
        right=Side(style="thin", color="D6CFC2"),
        top=Side(style="thin", color="D6CFC2"),
        bottom=Side(style="thin", color="D6CFC2"),
    )
    alt = PatternFill("solid", fgColor="F7F3EA")
    money = Font(name="Calibri")

    for col, h in enumerate(headers, 1):
        cell = ws.cell(1, col, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin

    for i, t in enumerate(txns, 2):
        values = [
            t.get("date"),
            t.get("description"),
            t.get("reference"),
            t.get("debit"),
            t.get("credit"),
            t.get("amount"),
            t.get("balance"),
            (t.get("type") or "").title(),
            t.get("suggested_ledger") or t.get("party_ledger"),
        ]
        for col, v in enumerate(values, 1):
            cell = ws.cell(i, col, v)
            cell.font = money
            cell.border = thin
            if i % 2 == 0:
                cell.fill = alt
            if col in (4, 5, 6, 7) and isinstance(v, (int, float)):
                cell.number_format = '#,##0.00'

    widths = [14, 48, 18, 14, 14, 14, 14, 12, 28]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = f"A1:I{max(1, len(txns) + 1)}"
    ws.freeze_panes = "A2"

    summary = wb.create_sheet("Summary")
    receipts = sum(t.get("credit") or 0 for t in txns)
    payments = sum(t.get("debit") or 0 for t in txns)
    summary["A1"] = "Metric"
    summary["B1"] = "Value"
    summary["A1"].font = header_font
    summary["B1"].font = header_font
    summary["A1"].fill = header_fill
    summary["B1"].fill = header_fill
    rows_s = [
        ("Transactions", len(txns)),
        ("Total credits (receipts)", receipts),
        ("Total debits (payments)", payments),
        ("Net", receipts - payments),
    ]
    for i, (k, v) in enumerate(rows_s, 2):
        summary.cell(i, 1, k)
        summary.cell(i, 2, v)
        if i > 2:
            summary.cell(i, 2).number_format = '#,##0.00'
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
